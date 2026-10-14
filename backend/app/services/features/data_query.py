"""Ask your data: natural language -> query plan -> deterministic local execution.

Small local models are unreliable at arithmetic and sorting. Instead of asking
the model for the *result*, we ask it for a *plan* (filters, sort, group-by,
aggregates, chart) constrained by a JSON schema, and execute that plan in
Python on the real table. Numbers are therefore always exact, and the plan is
shown to the user so the reasoning is transparent and auditable.
"""
from __future__ import annotations

import ast
import csv
import io
import operator
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any, Callable, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from app.models.document import DocumentContent
from app.services.parsers.csv_parser import sniff_delimiter
from app.services.parsers.text_parser import read_text_with_fallback
from app.utils.numbers import infer_decimal_comma, parse_number

from .structured import ask_structured

# ------------------------------------------------------------------ tables --


class Table(BaseModel):
    name: str
    columns: list[str]
    rows: list[list[Any]]
    sheets: list[str] = Field(default_factory=list)


class DataQueryError(Exception):
    pass


def _coerce(v: Any, decimal_comma: bool | None = None) -> Any:
    """A cell as a number when it is one ("184,6", "1 148", "45 %"), else as text."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "":
        return None
    number = parse_number(s, decimal_comma)
    return s if number is None else number


def _coerce_rows(rows: list[list[Any]], width: int, fallback_decimal_comma: bool | None) -> list[list[Any]]:
    """Pad rows to the header width and read numbers column by column.

    A column's own values decide whether its comma is a decimal mark ("184,6")
    or groups thousands ("1,148,000"); ``fallback_decimal_comma`` decides the
    columns that show neither (a semicolon-separated CSV is a European export).
    """
    padded = [list(r[:width]) + [None] * (width - len(r)) for r in rows]
    for j in range(width):
        cells = [r[j] for r in padded if isinstance(r[j], str)]
        convention = infer_decimal_comma(cells)
        convention = fallback_decimal_comma if convention is None else convention
        for r in padded:
            r[j] = _coerce(r[j], convention)
    return padded


def load_table(doc: DocumentContent, sheet: str | None = None) -> Table:
    """Read the real rows of a CSV/XLSX document (not the Markdown rendering)."""
    if not doc.source_path or not Path(doc.source_path).exists():
        raise DataQueryError("The original file for this document is no longer available.")
    path = Path(doc.source_path)
    if doc.source_type.value == "csv":
        text = read_text_with_fallback(path)
        delim = "\t" if path.suffix.lower() == ".tsv" else sniff_delimiter(text[:20000])
        rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]
        if not rows:
            raise DataQueryError("The CSV file is empty.")
        headers = [h.strip() or f"col{i + 1}" for i, h in enumerate(rows[0])]
        # A semicolon-separated export comes from a decimal-comma locale; in a
        # comma-separated file a quoted "1,148" groups thousands.
        fallback = {";": True, ",": False}.get(delim)
        data = _coerce_rows(rows[1:], len(headers), fallback_decimal_comma=fallback)
        return Table(name=doc.display_name, columns=headers, rows=data)
    if doc.source_type.value == "xlsx":
        wb = load_workbook(str(path), read_only=True, data_only=True)
        names = wb.sheetnames
        ws = wb[sheet] if sheet and sheet in names else wb[names[0]]
        raw = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        wb.close()
        if not raw:
            raise DataQueryError(f"Sheet '{ws.title}' is empty.")
        headers = [str(h).strip() if h not in (None, "") else f"col{i + 1}" for i, h in enumerate(raw[0])]
        data = _coerce_rows(raw[1:], len(headers), fallback_decimal_comma=None)
        return Table(name=f"{doc.display_name} / {ws.title}", columns=headers, rows=data, sheets=names)
    raise DataQueryError("Ask-your-data works with CSV and XLSX documents only.")


def describe_table(t: Table, sample_rows: int = 5) -> str:
    """Compact schema description for the model: columns, types, sample values."""
    lines = [f"Table: {t.name}", f"Rows: {len(t.rows)}", "Columns:"]
    for i, col in enumerate(t.columns):
        vals = [r[i] for r in t.rows if i < len(r) and r[i] is not None]
        nums = [v for v in vals if isinstance(v, (int, float))]
        kind = "number" if vals and len(nums) >= 0.8 * len(vals) else "text"
        distinct = list(OrderedDict.fromkeys(str(v) for v in vals))
        extra = ""
        if kind == "number" and nums:
            extra = f" min={min(nums)} max={max(nums)}"
        elif len(distinct) <= 12:
            extra = " values: " + ", ".join(distinct[:12])
        lines.append(f"- {col} ({kind}){extra}")
    lines.append("Sample rows:")
    for r in t.rows[:sample_rows]:
        lines.append("  " + " | ".join("" if v is None else str(v) for v in r))
    return "\n".join(lines)


# -------------------------------------------------------------------- plan --

Op = Literal["==", "!=", ">", ">=", "<", "<=", "contains", "not_contains", "starts_with", "in"]
Agg = Literal["count", "sum", "avg", "min", "max"]


class Filter(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: str
    op: Op
    value: str = Field(description="Comparison value as text; for 'in' use comma-separated values")


class Aggregate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: str = Field(description="Column to aggregate; for count any column")
    func: Agg
    alias: str = Field(default="", description="Output column name")


class Computed(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = Field(description="New column name")
    expression: str = Field(description="Arithmetic over existing numeric columns, e.g. (assignment_1 + assignment_2) / 2 or exam * 0.5")


class Sort(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: str
    descending: bool = True


class Chart(BaseModel):
    model_config = ConfigDict(extra="forbid")
    type: Literal["none", "bar", "line"] = "none"
    x: str = Field(default="", description="Category column")
    y: str = Field(default="", description="Numeric column")


class QueryPlan(BaseModel):
    """All fields are required on purpose: with grammar-constrained decoding the model must
    then make an explicit decision for every part of the plan instead of skipping it."""

    model_config = ConfigDict(extra="forbid")
    computed: list[Computed] = Field(description="New columns to compute first; [] if none")
    filters: list[Filter] = Field(description="Row filters; [] if none")
    group_by: str = Field(description="Column to group by, or empty string")
    aggregates: list[Aggregate] = Field(description="Aggregations; used with group_by, or alone for one overall total; [] if none")
    select: list[str] = Field(description="Columns to show; [] = all")
    sort: list[Sort] = Field(description="Sort order; [] if none")
    limit: int = Field(ge=0, description="Max rows; 0 = no limit")
    chart: Chart = Field(description="Chart to draw, type 'none' if not useful")
    explanation: str = Field(description="One sentence in plain language describing what the plan does")


PLAN_INSTRUCTIONS = """You translate a question about a table into a query plan. You must NOT compute any results yourself; the application executes the plan and shows the exact numbers.

{schema}

Question: {question}

Rules: use only column names that exist. Use 'computed' for derived values such as weighted averages. Use group_by + aggregates for totals per category. "Top N" / "highest" / "lowest" questions need a 'sort' entry AND a 'limit'. Use a chart when the question asks for a comparison or distribution. Keep 'explanation' to one plain sentence.

Example - question "top 3 products by revenue" on columns product, units, revenue:
{{"computed": [], "filters": [], "group_by": "", "aggregates": [], "select": ["product", "revenue"], "sort": [{{"column": "revenue", "descending": true}}], "limit": 3, "chart": {{"type": "bar", "x": "product", "y": "revenue"}}, "explanation": "Sort rows by revenue descending and keep the first three."}}

Example - question "average units per region" on columns region, product, units:
{{"computed": [], "filters": [], "group_by": "region", "aggregates": [{{"column": "units", "func": "avg", "alias": "avg_units"}}], "select": [], "sort": [{{"column": "avg_units", "descending": true}}], "limit": 0, "chart": {{"type": "bar", "x": "region", "y": "avg_units"}}, "explanation": "Group rows by region and average the units."}}"""


class QueryResult(BaseModel):
    columns: list[str]
    rows: list[list[Any]]
    row_count: int
    total_rows: int
    plan: QueryPlan
    table_name: str
    sheets: list[str] = Field(default_factory=list)


# --------------------------------------------------------------- executor --

# Computed columns are plain arithmetic, checked against a whitelist of syntax
# nodes and evaluated by walking the tree - never by eval(). Powers are left out
# on purpose: "9 ** 9 ** 9 ** 9" from a model-written plan would occupy the CPU
# for good.
_OPERATORS: dict[type, Callable[[Any, Any], Any]] = {
    ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv,
}
_ALLOWED_NODES = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Name, ast.Load,
                  ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd)
MAX_EXPRESSION = 300


def compile_expression(expr: str, columns: list[str]) -> Callable[[dict[str, Any]], Any]:
    """Check an arithmetic expression over the table's columns once; return a row evaluator.

    Column names may contain spaces ("Exam score * 0.5"): each is swapped for a
    placeholder where it stands as a whole name, longest first. Other names
    must resolve to a column (case and punctuation may differ).
    """
    if len(expr) > MAX_EXPRESSION:
        raise DataQueryError(f"The expression is too long ({len(expr)} characters).")
    names: dict[str, str] = {}
    text = expr
    for i, col in enumerate(sorted(columns, key=len, reverse=True)):
        placeholder = f"_col{i}"
        replaced = re.sub(rf"(?<!\w){re.escape(col)}(?!\w)", placeholder, text)
        if replaced != text:
            names[placeholder], text = col, replaced
    try:
        tree = ast.parse(text.strip(), mode="eval")
    except SyntaxError as exc:
        raise DataQueryError(f"Unsupported expression: {expr}") from exc
    for node in ast.walk(tree):
        if not isinstance(node, _ALLOWED_NODES) or (
            isinstance(node, ast.Constant) and (isinstance(node.value, bool) or not isinstance(node.value, (int, float)))
        ):
            raise DataQueryError(
                f"Unsupported expression: {expr}. Only numbers, column names, + - * / and brackets are allowed."
            )
        if isinstance(node, ast.Name) and node.id not in names:
            names[node.id] = resolve_column(node.id, columns, "expression")

    def evaluate(row: dict[str, Any]) -> Any:
        value = _evaluate(tree.body, row, names)
        return round(value, 4) if isinstance(value, float) else value

    return evaluate


def _evaluate(node: ast.AST, row: dict[str, Any], names: dict[str, str]) -> Any:
    """None propagates: a row whose input is missing or text gets no computed value."""
    if isinstance(node, ast.Constant):
        return node.value
    if isinstance(node, ast.Name):
        value = row.get(names[node.id])
        return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None
    if isinstance(node, ast.UnaryOp):
        value = _evaluate(node.operand, row, names)
        return None if value is None else (-value if isinstance(node.op, ast.USub) else value)
    left, right = _evaluate(node.left, row, names), _evaluate(node.right, row, names)  # type: ignore[attr-defined]
    if left is None or right is None:
        return None
    try:
        return _OPERATORS[type(node.op)](left, right)  # type: ignore[attr-defined]
    except ZeroDivisionError:
        return None


def _cmp_value(v: Any, target: str) -> Any:
    t = _coerce(target)
    if isinstance(v, (int, float)) and isinstance(t, (int, float)):
        return v, t
    return ("" if v is None else str(v)).lower(), str(target).lower()


def _match(row: dict[str, Any], f: Filter) -> bool:
    if f.column not in row:
        raise DataQueryError(f"Unknown column in filter: {f.column}")
    v = row[f.column]
    if f.op == "in":
        wanted = [w.strip().lower() for w in f.value.split(",")]
        return ("" if v is None else str(v)).lower() in wanted
    a, b = _cmp_value(v, f.value)
    if f.op in ("contains", "not_contains", "starts_with"):
        a, b = str(a), str(b)
        hit = b in a if f.op != "starts_with" else a.startswith(b)
        return not hit if f.op == "not_contains" else hit
    if v is None:
        return f.op == "!="
    try:
        return {"==": a == b, "!=": a != b, ">": a > b, ">=": a >= b, "<": a < b, "<=": a <= b}[f.op]
    except TypeError:
        return False


def _key(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def resolve_column(name: str, cols: list[str], what: str) -> str:
    """Exact match first, then case/punctuation-insensitive, then unique substring."""
    if name in cols:
        return name
    k = _key(name)
    for c in cols:
        if _key(c) == k:
            return c
    partial = [c for c in cols if k and (k in _key(c) or _key(c) in k)]
    if len(partial) == 1:
        return partial[0]
    raise DataQueryError(f"Unknown column in {what}: '{name}'. Available: {', '.join(cols)}")


def _normalise_plan(table: Table, plan: QueryPlan) -> None:
    """Resolve every column reference in the plan (models vary spelling and case)."""
    cols = list(table.columns) + [c.name for c in plan.computed if c.name not in table.columns]
    if plan.group_by and not plan.aggregates:
        # "group by student" with nothing to aggregate is just a per-row listing
        plan.group_by = ""
    for f in plan.filters:
        f.column = resolve_column(f.column, cols, "filter")
    if plan.group_by:
        plan.group_by = resolve_column(plan.group_by, cols, "group_by")
    for a in plan.aggregates:
        if a.func != "count" or a.column:
            a.column = resolve_column(a.column, cols, "aggregate")
    if plan.group_by:
        # grouped output = key + aggregates; other columns stay reachable (first value in the group)
        out_cols = [plan.group_by] + [a.alias or f"{a.func}_{a.column}" for a in plan.aggregates] + [c for c in cols if c != plan.group_by]
    elif plan.aggregates:
        out_cols = [a.alias or f"{a.func}_{a.column}" for a in plan.aggregates]
    else:
        out_cols = cols
    for srt in plan.sort:
        srt.column = resolve_column(srt.column, out_cols, "sort")
    plan.select = [resolve_column(c, out_cols, "select") for c in plan.select]
    if plan.chart.type != "none":
        try:
            plan.chart.x = resolve_column(plan.chart.x, out_cols, "chart")
            plan.chart.y = resolve_column(plan.chart.y, out_cols, "chart")
        except DataQueryError:
            plan.chart = Chart()


def execute_plan(table: Table, plan: QueryPlan) -> QueryResult:
    _normalise_plan(table, plan)
    cols = list(table.columns)
    rows: list[dict[str, Any]] = [dict(zip(cols, r)) for r in table.rows]

    for c in plan.computed:
        evaluate = compile_expression(c.expression, cols)
        for r in rows:
            r[c.name] = evaluate(r)
        if c.name not in cols:
            cols.append(c.name)

    for f in plan.filters:
        rows = [r for r in rows if _match(r, f)]

    if plan.group_by:
        groups: "OrderedDict[Any, list[dict[str, Any]]]" = OrderedDict()
        for r in rows:
            groups.setdefault(r[plan.group_by], []).append(r)
        out_cols = [plan.group_by]
        aggs = plan.aggregates or [Aggregate(column=plan.group_by, func="count", alias="count")]
        out_rows = []
        for key, members in groups.items():
            out = dict(members[0])  # any-value semantics for non-aggregated columns
            out[plan.group_by] = key
            for a in aggs:
                name = a.alias or f"{a.func}_{a.column}"
                out[name] = _aggregate(members, a)
                if name not in out_cols:
                    out_cols.append(name)
            out_rows.append(out)
        rows, cols = out_rows, out_cols
    elif plan.aggregates:
        out = {}
        for a in plan.aggregates:
            name = a.alias or f"{a.func}_{a.column}"
            out[name] = _aggregate(rows, a)
        rows, cols = [out], list(out.keys())

    for s in reversed(plan.sort):
        rows.sort(key=lambda r: _sort_key(r.get(s.column)), reverse=s.descending)

    if plan.select:
        cols = list(plan.select)

    total = len(rows)
    if plan.limit:
        rows = rows[: plan.limit]
    return QueryResult(
        columns=cols,
        rows=[[r.get(c) for c in cols] for r in rows],
        row_count=len(rows),
        total_rows=total,
        plan=plan,
        table_name=table.name,
        sheets=table.sheets,
    )


def _sort_key(v: Any):
    if v is None:
        return (0, 0, "")
    if isinstance(v, (int, float)):
        return (1, v, "")
    return (2, 0, str(v).lower())


def _aggregate(rows: list[dict[str, Any]], a: Aggregate) -> Any:
    if a.func == "count":
        return len(rows)
    vals = [r.get(a.column) for r in rows]
    nums = [v for v in vals if isinstance(v, (int, float))]
    if not nums:
        return None
    if a.func == "sum":
        return round(sum(nums), 4)
    if a.func == "avg":
        return round(sum(nums) / len(nums), 4)
    if a.func == "min":
        return min(nums)
    return max(nums)


async def ask_data(doc: DocumentContent, question: str, sheet: str | None = None) -> QueryResult:
    table = load_table(doc, sheet)
    plan = await ask_structured(
        QueryPlan,
        PLAN_INSTRUCTIONS.format(schema=describe_table(table), question=question.strip()),
        None,
        system_prompt="You are a careful data analyst who only writes query plans as JSON.",
        what="query plan",
        max_tokens=1200,
    )
    return execute_plan(table, plan)
