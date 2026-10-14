"""PPTX parser using python-pptx.

Every slide becomes a section "Slide N" with the slide title, body text
(bullets preserved by paragraph level), tables, charts and speaker notes.
Images are listed by count only; they are not interpreted here (the Figures
tab reads them with a vision model).

Native charts (chart objects, not pictures of charts) keep their data inside
the file, so their categories and values are written out as a table. Before,
a slide holding only a chart came out as "no text on this slide" and its
numbers were invisible to the model.
"""
from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from app.models.document import DocumentSection, SourceType

from .base import BaseParser, ParseError, make_section, markdown_table

log = logging.getLogger(__name__)


def _number(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _chart_md(chart) -> str:
    """The data of a native chart as Markdown: one row per category, one column per series."""
    title = ""
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = chart.chart_title.text_frame.text.strip()
    except Exception:  # noqa: BLE001 - a missing title is fine
        pass
    kind = getattr(chart.chart_type, "name", "") or ""
    tables: list[str] = []
    for plot in chart.plots:
        series = list(plot.series)
        if not series:
            continue
        try:
            categories = [str(c) for c in plot.categories]
        except Exception:  # noqa: BLE001 - e.g. XY charts have no category axis
            categories = []
        first_column = "Category"
        if not categories:
            # A scatter chart has x values instead of categories.
            xs = series[0]._element.xpath("./c:xVal//c:pt/c:v")
            if xs:
                categories, first_column = [_number(float(x.text)) for x in xs], "X"
        length = max(len(categories), *(len(s.values) for s in series))
        headers = [first_column] + [s.name or f"Series {i + 1}" for i, s in enumerate(series)]
        rows = [
            [categories[k] if k < len(categories) else str(k + 1)]
            + [_number(s.values[k]) if k < len(s.values) else "" for s in series]
            for k in range(length)
        ]
        tables.append(markdown_table(headers, rows))
    if not tables:
        return ""
    head = f"### Chart: {title}" if title else "### Chart"
    kind_note = f": {kind.replace('_', ' ').lower()}" if kind else ""
    return f"{head}\n\n" + "\n\n".join(tables) + f"\n\n_(chart data read from the file{kind_note})_"


def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _text_frame_md(tf, is_title: bool = False) -> list[str]:
    lines: list[str] = []
    for para in tf.paragraphs:
        text = "".join(r.text for r in para.runs).strip() or para.text.strip()
        if not text:
            continue
        if is_title:
            lines.append(text)
        else:
            indent = "  " * max(para.level or 0, 0)
            lines.append(f"{indent}* {text}")
    return lines


class PptxParser(BaseParser):
    source_type = SourceType.pptx
    extensions = ("pptx",)

    def parse(self, path: Path, display_name: str) -> list[DocumentSection]:
        try:
            prs = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            raise ParseError("The PowerPoint file could not be opened. Is it a valid .pptx file?") from exc

        sections: list[DocumentSection] = []
        for idx, slide in enumerate(prs.slides, start=1):
            title = ""
            body: list[str] = []
            tables: list[str] = []
            charts: list[str] = []
            image_count = 0
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame:
                title = title_shape.text_frame.text.strip().replace("\n", " ")

            # sort shapes top-to-bottom, left-to-right for a sensible reading order
            shapes = sorted(
                _iter_shapes(slide.shapes),
                key=lambda s: (Emu(s.top or 0), Emu(s.left or 0)),
            )
            for sh in shapes:
                if title_shape is not None and sh.shape_id == title_shape.shape_id:
                    continue
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
                    continue
                if getattr(sh, "has_table", False) and sh.has_table:
                    rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                    if rows:
                        tables.append(markdown_table(rows[0], rows[1:]))
                    continue
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    try:
                        md = _chart_md(sh.chart)
                    except Exception as exc:  # noqa: BLE001 - one odd chart must not fail the file
                        log.info("Chart on slide %d of %s could not be read: %s", idx, display_name, exc)
                        md = ""
                    if md:
                        charts.append(md)
                    continue
                if sh.has_text_frame:
                    body.extend(_text_frame_md(sh.text_frame))

            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()

            parts: list[str] = []
            if body:
                parts.append("\n".join(body))
            for t in tables:
                parts.append("### Table\n\n" + t)
            parts.extend(charts)
            if notes:
                parts.append("### Speaker notes\n\n" + notes)
            if image_count:
                parts.append(f"_({image_count} image(s) on this slide, not interpreted)_")
            md = "\n\n".join(parts).strip() or "_(no text on this slide)_"
            sections.append(
                make_section(
                    len(sections) + 1,
                    title or f"Slide {idx}",
                    f"Slide {idx}",
                    md,
                    slide=idx,
                    images=image_count,
                    charts=len(charts),
                )
            )
        if not sections:
            raise ParseError("The presentation contains no slides.")
        return sections

    def to_document(self, path, original_filename, display_name=None):
        doc = super().to_document(path, original_filename, display_name)
        doc.metadata["slide_count"] = len(doc.sections)
        return doc
