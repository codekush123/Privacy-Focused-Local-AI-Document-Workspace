"""XLSX parser using openpyxl.

Workbooks are opened read-only with ``data_only=True`` so cached formula
results are used and formulas are never evaluated. Macros are never executed
(openpyxl does not run VBA). Each worksheet becomes one section
("Sheet: <name>") containing a Markdown table of the used range.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.models.document import DocumentSection, SourceType

from .base import BaseParser, ParseError, make_section, markdown_table

MAX_ROWS_PER_SHEET = 20000  # hard safety cap to avoid freezing on gigantic sheets


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        try:
            return v.isoformat(sep=" ") if hasattr(v, "hour") else v.isoformat()
        except TypeError:
            return v.isoformat()
    return str(v)


class XlsxParser(BaseParser):
    source_type = SourceType.xlsx
    extensions = ("xlsx", "xlsm")

    def parse(self, path: Path, display_name: str) -> list[DocumentSection]:
        try:
            wb = load_workbook(str(path), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ParseError("The Excel workbook could not be opened. Is it a valid .xlsx file?") from exc
        # A second pass (not data_only) is needed to detect formulas; keep it cheap.
        formula_cells = 0
        try:
            wb_f = load_workbook(str(path), read_only=True, data_only=False)
            for ws in wb_f.worksheets:
                for row in ws.iter_rows(values_only=True, max_row=min(ws.max_row or 0, 2000)):
                    formula_cells += sum(1 for v in row if isinstance(v, str) and v.startswith("="))
            wb_f.close()
        except Exception:  # noqa: BLE001
            pass

        sections: list[DocumentSection] = []
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            truncated = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS_PER_SHEET:
                    truncated = True
                    break
                cells = [_fmt(v) for v in row]
                if any(c.strip() for c in cells):
                    rows.append(cells)
            # drop fully empty trailing columns
            width = 0
            for r in rows:
                for j in range(len(r) - 1, -1, -1):
                    if r[j].strip():
                        width = max(width, j + 1)
                        break
            rows = [r[:width] for r in rows]
            if not rows:
                md = "_(empty sheet)_"
            else:
                md = markdown_table(rows[0], rows[1:])
                if truncated:
                    md += f"\n\n_(sheet truncated at {MAX_ROWS_PER_SHEET} rows for safety - this is reported, not hidden)_"
            sections.append(
                make_section(
                    len(sections) + 1,
                    ws.title,
                    f"Sheet: {ws.title}",
                    md,
                    sheet=ws.title,
                    row_count=max(len(rows) - 1, 0),
                    column_count=width,
                    truncated=truncated,
                )
            )
        wb.close()
        if not sections:
            raise ParseError("The workbook contains no worksheets.")
        sections[0].metadata["formula_cells"] = formula_cells
        return sections

    def to_document(self, path, original_filename, display_name=None):
        doc = super().to_document(path, original_filename, display_name)
        doc.metadata["sheets"] = [s.title for s in doc.sections]
        fc = doc.sections[0].metadata.get("formula_cells", 0) if doc.sections else 0
        if fc:
            doc.metadata["formula_cells"] = fc
            doc.metadata["note"] = (
                "Formulas were not evaluated; cached values stored in the file were used where available."
            )
        return doc
