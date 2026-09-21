"""Word (.docx) writer built on python-docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.schemas.artifacts import DocxSpec


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    ncols = max(len(headers), max((len(r) for r in rows), default=0))
    if ncols == 0:
        return
    table = doc.add_table(rows=1 if headers else 0, cols=ncols)
    table.style = "Light Grid Accent 1"
    if headers:
        hdr = table.rows[0].cells
        for i in range(ncols):
            hdr[i].text = headers[i] if i < len(headers) else ""
            for p in hdr[i].paragraphs:
                for r in p.runs:
                    r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i in range(ncols):
            cells[i].text = str(row[i]) if i < len(row) else ""
    doc.add_paragraph()


def write_docx(spec: DocxSpec, out_path: Path, sources: list[str] | None = None) -> Path:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    doc.add_heading(spec.title, level=0)
    if spec.subtitle:
        p = doc.add_paragraph(spec.subtitle)
        p.runs[0].italic = True

    for block in spec.blocks:
        if block.type == "heading":
            doc.add_heading(block.text or " ", level=min(max(block.level, 1), 3))
        elif block.type == "paragraph":
            if block.text:
                doc.add_paragraph(block.text)
        elif block.type == "bullets":
            for item in block.items:
                doc.add_paragraph(item, style="List Bullet")
        elif block.type == "numbered":
            for item in block.items:
                doc.add_paragraph(item, style="List Number")
        elif block.type == "table":
            _add_table(doc, block.headers, block.rows)

    if sources:
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run("Sources: " + ", ".join(sources))
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        # Footer as well
        footer = doc.sections[0].footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.text = "Generated locally from: " + ", ".join(sources)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in fp.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
