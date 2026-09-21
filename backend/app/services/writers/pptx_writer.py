"""PowerPoint (.pptx) writer built on python-pptx."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.schemas.artifacts import PptxSpec

MAX_BULLETS_PER_SLIDE = 8


def _add_source_footer(slide, prs, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.3), prs.slide_height - Inches(0.45), prs.slide_width - Inches(0.6), Inches(0.35))
    tf = box.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def write_pptx(spec: PptxSpec, out_path: Path, sources: list[str] | None = None) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]
    footer_text = ("Sources: " + ", ".join(sources)) if sources else ""

    # Title slide
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = spec.presentation_title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = spec.subtitle or "Generated locally from selected documents"
    if footer_text:
        _add_source_footer(slide, prs, footer_text)

    for s in spec.slides:
        bullets = [b for b in s.bullet_points if b and b.strip()]
        # Split long bullet lists over continuation slides so text stays readable.
        chunks = [bullets[i : i + MAX_BULLETS_PER_SLIDE] for i in range(0, len(bullets), MAX_BULLETS_PER_SLIDE)] or [[]]
        for ci, chunk in enumerate(chunks):
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = s.title if ci == 0 else f"{s.title} (cont.)"
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.word_wrap = True
            if chunk:
                tf.text = chunk[0]
                for b in chunk[1:]:
                    p = tf.add_paragraph()
                    p.text = b
                for p in tf.paragraphs:
                    p.font.size = Pt(20 if len(chunk) <= 5 else 16)
            else:
                tf.text = ""
            if s.notes and ci == 0:
                slide.notes_slide.notes_text_frame.text = s.notes

    if sources:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = "Sources"
        tf = slide.placeholders[1].text_frame
        tf.text = sources[0]
        for src in sources[1:]:
            tf.add_paragraph().text = src
        for p in tf.paragraphs:
            p.font.size = Pt(18)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
