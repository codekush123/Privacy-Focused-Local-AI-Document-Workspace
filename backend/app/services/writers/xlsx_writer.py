"""Excel (.xlsx) writer built on openpyxl.

Every cell value comes from the model or from a document, so none of it may
become a live formula: openpyxl stores any string that starts with "=" as one,
and a formula such as =HYPERLINK("https://…?"&B2) or =WEBSERVICE(…) in a
generated workbook would send its contents off the machine when the file is
opened in Excel. Such text is written as text.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.artifacts import XlsxSpec
from app.utils.numbers import infer_decimal_comma, parse_number

_INVALID_SHEET_CHARS = re.compile(r"[\[\]\*\?/\\:]")


def _sheet_name(name: str, used: set[str], fallback: str) -> str:
    base = _INVALID_SHEET_CHARS.sub(" ", name or fallback).strip()[:31] or fallback
    candidate, n = base, 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _coerce(value: str, decimal_comma: bool | None = None):
    """Store numbers as numbers so Excel can compute on them: "184.6", and also
    the Finnish "184,6" and "1 148". Percentages, codes with leading zeros and
    very long digit strings (account numbers) stay text."""
    if value is None:
        return ""
    s = str(value).strip()
    if s.endswith("%") or len(re.sub(r"\D", "", s)) >= 16:
        return s
    number = parse_number(s, decimal_comma)
    return s if number is None else number


def _append(ws, values: list) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        if cell.data_type == "f":  # text that merely starts with "=" - never a formula
            cell.data_type = "s"


def write_xlsx(spec: XlsxSpec, out_path: Path, sources: list[str] | None = None) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")

    for idx, sheet in enumerate(spec.sheets, start=1):
        ws = wb.create_sheet(_sheet_name(sheet.name, used, f"Sheet{idx}"))
        ncols = max(len(sheet.headers), max((len(r) for r in sheet.rows), default=0))
        if sheet.headers:
            _append(ws, list(sheet.headers) + [""] * (ncols - len(sheet.headers)))
            for c in ws[1]:
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(vertical="top", wrap_text=True)
            ws.freeze_panes = "A2"
        # Each column's own values decide whether its comma is a decimal mark.
        conventions = [infer_decimal_comma([r[j] for r in sheet.rows if j < len(r)]) for j in range(ncols)]
        for row in sheet.rows:
            _append(ws, [_coerce(v, conventions[j]) for j, v in enumerate(row)] + [""] * (ncols - len(row)))
        # column widths: fit to content, capped
        for col in range(1, ncols + 1):
            longest = 8
            for cell in ws[get_column_letter(col)]:
                if cell.value is not None:
                    longest = max(longest, min(len(str(cell.value)), 60))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = longest + 2
        if sheet.notes:
            ws.append([])
            _append(ws, [sheet.notes])
            ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="666666")
        if ncols:
            ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(len(sheet.rows) + 1, 1)}"

    if not wb.worksheets:
        wb.create_sheet("Sheet1")

    if sources:
        ws = wb.create_sheet("Sources")
        ws.append(["Generated locally from"])
        ws["A1"].font = Font(bold=True)
        for s in sources:
            _append(ws, [s])
        ws.column_dimensions["A"].width = 60

    wb.properties.title = spec.workbook_title
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path
