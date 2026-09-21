"""Generate small fixture files for parser tests (run: python tests/make_fixtures.py)."""
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

FIX = Path(__file__).parent / "fixtures"
FIX.mkdir(exist_ok=True)

MULTI = "English hello | Finnish: Hyvää päivää, ääkköset | Chinese: 你好，世界 | Arabic: مرحبا بالعالم | Russian: Привет, мир"

(FIX / "sample.txt").write_text("Plain text fixture.\nThe verification phrase is BLUE ELEPHANT 1947.\n" + MULTI + "\n", encoding="utf-8")
(FIX / "sample.md").write_text("# Markdown Title\n\nSome *markdown* text.\n\nThe verification phrase is BLUE ELEPHANT 1947.\n\n- item one\n- item two\n\n" + MULTI + "\n", encoding="utf-8")
(FIX / "sample.html").write_text(
    """<!doctype html><html><head><title>Fixture Page</title><style>body{color:red}</style>
<script>alert('should not appear SCRIPT_MARKER')</script></head><body>
<nav><a href="/">NAV_MARKER home</a></nav>
<main><h1>HTML Heading</h1><p>The verification phrase is BLUE ELEPHANT 1947.</p>
<ul><li>alpha</li><li>beta</li></ul>
<table><tr><th>Name</th><th>Score</th></tr><tr><td>Ada</td><td>95</td></tr></table>
<p style="display:none">HIDDEN_MARKER</p><p>""" + MULTI + """</p></main>
<footer>FOOTER_MARKER</footer></body></html>""",
    encoding="utf-8",
)
(FIX / "sample.csv").write_text("Name,Score,Grade\nAda,95,A\nBob,71,C\n\"Smith, Jr.\",88,B\nPhrase,0,BLUE ELEPHANT 1947\nUnicode,100,\"" + MULTI + "\"\n", encoding="utf-8")
(FIX / "semicolon.csv").write_text("id;value\n1;BLUE ELEPHANT 1947\n2;other\n", encoding="utf-8")

d = Document()
d.add_heading("Introduction", level=1)
d.add_paragraph("The verification phrase is BLUE ELEPHANT 1947.")
d.add_paragraph("First bullet", style="List Bullet")
d.add_paragraph("Second bullet", style="List Bullet")
d.add_heading("Details", level=2)
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text = "Key"; t.cell(0, 1).text = "Value"
t.cell(1, 0).text = "Color"; t.cell(1, 1).text = "TABLE_CELL_VALUE"
d.add_paragraph(MULTI)
d.save(FIX / "sample.docx")

pdf = pymupdf.open()
p = pdf.new_page(); p.insert_text((72, 72), "Page one content.", fontsize=12)
p = pdf.new_page(); p.insert_text((72, 72), "Page two content.", fontsize=12)
p = pdf.new_page(); p.insert_text((72, 72), "The verification phrase is BLUE ELEPHANT 1947.", fontsize=12)
pdf.save(FIX / "sample.pdf"); pdf.close()

wb = Workbook()
ws = wb.active; ws.title = "Students"
ws.append(["Name", "Score", "Grade"]); ws.append(["Ada", 95, "A"]); ws.append(["Bob", 71, "C"])
ws2 = wb.create_sheet("Results")
ws2.append(["Key", "Value"]); ws2.append(["Phrase", "BLUE ELEPHANT 1947"]); ws2.append(["Multi", MULTI])
wb.save(FIX / "sample.xlsx")

prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Intro Slide"; s.placeholders[1].text = "First slide bullet"
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Second Slide"
tf = s.placeholders[1].text_frame; tf.text = "The verification phrase is BLUE ELEPHANT 1947."
p2 = tf.add_paragraph(); p2.text = MULTI; p2.level = 1
rows, cols = 2, 2
shape = s.shapes.add_table(rows, cols, Inches(1), Inches(4), Inches(6), Inches(1))
shape.table.cell(0, 0).text = "H1"; shape.table.cell(0, 1).text = "H2"
shape.table.cell(1, 0).text = "PPTX_TABLE_CELL"; shape.table.cell(1, 1).text = "x"
prs.save(FIX / "sample.pptx")
print("fixtures written to", FIX)
