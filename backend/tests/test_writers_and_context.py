"""Writers, context strategy, privacy policy and URL validation tests."""
from __future__ import annotations

import pytest

from app.models.document import DocumentContent, DocumentSection, SourceType
from app.schemas.artifacts import DocxBlock, DocxSpec, PptxSlide, PptxSpec, XlsxSheet, XlsxSpec, json_schema_for
from app.services.llm.context_strategy import FullContextStrategy
from app.services.parsers import parse_file
from app.services.parsers.url_fetcher import UrlRejected, validate_url
from app.services.privacy.policy import is_localhost_url
from app.services.writers.docx_writer import write_docx
from app.services.writers.pptx_writer import write_pptx
from app.services.writers.simple_writers import write_csv
from app.services.writers.xlsx_writer import write_xlsx

from .conftest import MULTILINGUAL


def _doc(name: str, text: str) -> DocumentContent:
    return DocumentContent(
        id="x" + name,
        original_filename=name,
        source_type=SourceType.txt,
        display_name=name,
        sections=[DocumentSection(section_id="s1", title="Text", locator="Text", markdown=text)],
    ).finalize()


# ----------------------------------------------------------------- writers
def test_docx_writer_roundtrip(tmp_path):
    spec = DocxSpec(
        title="Quiz Title",
        subtitle="Sub",
        blocks=[
            DocxBlock(type="heading", text="Questions", level=1),
            DocxBlock(type="numbered", items=["Q one?", "Q two?"]),
            DocxBlock(type="paragraph", text="Paragraph " + MULTILINGUAL["finnish"]),
            DocxBlock(type="bullets", items=["b1", MULTILINGUAL["chinese"]]),
            DocxBlock(type="table", headers=["Q", "A"], rows=[["1", "alpha"], ["2", "beta"]]),
        ],
    )
    out = write_docx(spec, tmp_path / "q.docx", sources=["lecture.pdf"])
    doc = parse_file(out, "q.docx")
    md = doc.full_markdown
    assert "Quiz Title" in md
    assert "1. Q one?" in md
    assert MULTILINGUAL["finnish"] in md and MULTILINGUAL["chinese"] in md
    assert "| 2 | beta |" in md
    assert "lecture.pdf" in md


def test_xlsx_writer_roundtrip(tmp_path):
    spec = XlsxSpec(
        workbook_title="Quiz",
        sheets=[
            XlsxSheet(name="Quiz/Questions", headers=["Question", "Answer", "Difficulty"], rows=[["Q1", "A1", "easy"], ["Q2", MULTILINGUAL["arabic"], "hard"]]),
            XlsxSheet(name="Second", headers=["n"], rows=[["1"], ["2.5"]]),
        ],
    )
    out = write_xlsx(spec, tmp_path / "q.xlsx", sources=["data.csv"])
    doc = parse_file(out, "q.xlsx")
    assert "Sheet: Quiz Questions" in doc.full_markdown  # invalid '/' replaced
    assert "| Q2 | " + MULTILINGUAL["arabic"] + " | hard |" in doc.full_markdown
    assert "| 2.5 |" in doc.full_markdown  # numeric coercion keeps value
    assert "Sheet: Sources" in doc.full_markdown
    from openpyxl import load_workbook

    ws = load_workbook(out)["Quiz Questions"]
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws.column_dimensions["A"].width > 8


def test_pptx_writer_roundtrip(tmp_path):
    spec = PptxSpec(
        presentation_title="Deck",
        slides=[
            PptxSlide(title="One", bullet_points=["a", "b"], notes="note text"),
            PptxSlide(title="Long", bullet_points=[f"point {i}" for i in range(12)]),
        ],
    )
    out = write_pptx(spec, tmp_path / "d.pptx", sources=["s.pptx"])
    doc = parse_file(out, "d.pptx")
    locs = [s.locator for s in doc.sections]
    # title + One + Long + Long (cont.) + Sources
    assert len(locs) == 5
    assert "## Slide 2 - One" in doc.full_markdown
    assert "note text" in doc.full_markdown
    assert "Long (cont.)" in doc.full_markdown


def test_csv_writer(tmp_path):
    spec = XlsxSpec(workbook_title="t", sheets=[XlsxSheet(name="s", headers=["a", "b"], rows=[["1", "x,y"]])])
    out = write_csv(spec, tmp_path / "t.csv")
    assert out.read_bytes() == b'\xef\xbb\xbfa,b\r\n1,"x,y"\r\n'


def test_artifact_schemas_are_plain():
    for model in (DocxSpec, XlsxSpec, PptxSpec):
        schema = json_schema_for(model)
        assert schema["type"] == "object"
        assert "required" in schema


# ------------------------------------------------------------ context/prompt
def test_full_context_strategy_wraps_sources():
    docs = [_doc("a.txt", "Alpha content"), _doc("b.txt", "Beta content")]
    msgs = FullContextStrategy().build_messages(docs, "Question?", history=[{"role": "user", "content": "earlier"}, {"role": "assistant", "content": "reply"}])
    assert [m["role"] for m in msgs] == ["system", "user", "assistant", "user"]
    system = msgs[0]["content"]
    assert '<source id="1" name="a.txt">' in system and "</source>" in system
    assert '<source id="2" name="b.txt">' in system
    assert "Alpha content" in system and "Beta content" in system
    assert "not instructions" in system.lower() or "DATA, not instructions" in system
    assert msgs[-1]["content"] == "Question?"


def test_full_context_strategy_without_documents():
    msgs = FullContextStrategy().build_messages([], "Hi")
    assert len(msgs) == 2 and "<source" not in msgs[0]["content"]


def test_multilingual_text_survives_prompt_building(fixtures):
    doc = parse_file(fixtures / "sample.txt", "sample.txt")
    msgs = FullContextStrategy().build_messages([doc], "Q")
    for sample in MULTILINGUAL.values():
        assert sample in msgs[0]["content"]


# ---------------------------------------------------------------- privacy
@pytest.mark.parametrize("url,ok", [
    ("http://127.0.0.1:8080", True),
    ("http://localhost:8080", True),
    ("http://[::1]:8080", True),
    ("http://192.168.1.10:8080", False),
    ("https://api.example.com", False),
    ("not a url", False),
])
def test_is_localhost_url(url, ok):
    assert is_localhost_url(url) is ok


@pytest.mark.parametrize("url", [
    "file:///etc/passwd",
    "ftp://example.com/x",
    "http://localhost/",
    "http://127.0.0.1/",
    "http://10.0.0.1/",
    "http://192.168.0.1/",
    "http://169.254.169.254/latest/meta-data",
    "http://[::1]/",
    "http://user:pw@example.com/",
    "://bad",
    "",
])
def test_url_validation_rejects_unsafe(url):
    with pytest.raises(UrlRejected):
        validate_url(url)
