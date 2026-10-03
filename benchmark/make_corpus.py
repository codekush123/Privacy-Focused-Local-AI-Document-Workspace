"""Write the benchmark corpus (PDF, DOCX, PPTX, XLSX in English and Finnish).

Run from the repository root with the backend environment:
    backend/.venv/Scripts/python benchmark/make_corpus.py

The files are committed, so this is only needed after editing corpus_text.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from corpus_text import MINUTES, REGISTER, REPORT, SLIDES  # noqa: E402

OUT = HERE / "corpus"


def write_pdf(spec: dict, path: Path) -> None:
    doc = pymupdf.open()
    for i, (heading, paragraphs) in enumerate(spec["pages"]):
        page = doc.new_page(width=595, height=842)  # A4
        html = ""
        if i == 0:
            html += f"<h1>{spec['title']}</h1>"
        html += f"<h2>{heading}</h2>" + "".join(f"<p>{p}</p>" for p in paragraphs)
        page.insert_htmlbox(pymupdf.Rect(60, 60, 535, 790), html, css="* {font-family: sans-serif; font-size: 10.5pt;}")
    doc.set_metadata({"title": spec["title"], "author": "Kuusiranta Energy (fictional)"})
    doc.save(path)


def write_docx(spec: dict, path: Path) -> None:
    doc = Document()
    doc.add_heading(spec["title"], level=1)
    for heading, paragraphs, table in spec["sections"]:
        doc.add_heading(heading, level=2)
        for p in paragraphs:
            doc.add_paragraph(p)
        if table:
            t = doc.add_table(rows=len(table), cols=len(table[0]))
            t.style = "Table Grid"
            for r, row in enumerate(table):
                for c, value in enumerate(row):
                    t.cell(r, c).text = str(value)
    doc.save(path)


def write_pptx(spec: dict, path: Path) -> None:
    prs = Presentation()
    for i, (title, bullets) in enumerate(spec["slides"]):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        for para in body.paragraphs:
            for run in para.runs:
                run.font.size = Pt(20)
    prs.save(path)


def write_xlsx(spec: dict, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in spec["sheets"].items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value)) for c in col) + 2
    wb.save(path)


def main() -> None:
    for lang in ("en", "fi"):
        folder = OUT / lang
        folder.mkdir(parents=True, exist_ok=True)
        write_pdf(REPORT[lang], folder / REPORT[lang]["filename"])
        write_docx(MINUTES[lang], folder / MINUTES[lang]["filename"])
        write_pptx(SLIDES[lang], folder / SLIDES[lang]["filename"])
        write_xlsx(REGISTER[lang], folder / REGISTER[lang]["filename"])
        print(f"{lang}: wrote {len(list(folder.iterdir()))} files to {folder}")


if __name__ == "__main__":
    main()
